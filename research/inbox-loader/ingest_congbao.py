#!/usr/bin/env python3
"""Turn documents downloaded from Công báo into the canonical legal extracts.

    python3 ingest_congbao.py --work <dir-with-downloads> [--dry-run] [--only <số hiệu> ...] [--out <dir>]

Writes (idempotently — rows for the same document number are replaced, others kept; --out
writes the same four files into another directory, e.g. a scratch preview):
    db/seed/data/legal/documents.ndjson     one row per instrument, verification=auto_unverified
    db/seed/data/legal/provisions.ndjson    Chương/Mục/Điều/Khoản/Điểm tree
    db/seed/data/legal/chunks.ndjson        one chunk per Điều, embedded at seed time
    db/seed/data/legal/annex-tables.ndjson  every table, cell by cell (not yet seeded)

These files ARE the backup: `yarn db:seed:legal` loads the first three, so nothing done
while the project has no server has to be redone when it gets one back.

TWO BRANCHES, BECAUSE ONE WOULD LOSE THE ANNEXES. The project's prose parser
(research/legal-loader/parse_provisions.py) drops every line carrying a Word cell mark —
right for prose, fatal for a danh mục: NĐ 292/2026 alone has 66 tables, 349 rows. So
prose goes through parse_provisions.build() verbatim, and tables go through a table
reader that keeps cells. The gate refuses a document whose Điều do not run 1..N.

WORD FILES FROM CÔNG BÁO NEED TWO REPAIRS, both found the hard way (2026-09-10):
  * .docx zip entries use Windows back-slashes (`customXml\\item1.xml`) and include a
    `[trash]` entry; neither textutil nor python-docx opens them until repacked.
  * a heading paragraph can hold a soft line break ("Mục 2\\nĐÁNH GIÁ SỰ PHÙ HỢP"). Fed
    as one line, the parser takes the NEXT paragraph — "Điều 71. …" — for the wrapped
    tail of the section title and swallows the article. NĐ 37/2026 lost 5 of 99 that way.

TABLE CELLS COME FROM THE RAW w:tc, NOT python-docx `row.cells`, which repeats a merged cell in
every grid column and row it covers: TT 125/2026/TT-BCA's group name "Pháo hoa" landed in the
Mã HS column, and its one cell of 13 codes (8525.81.10–8525.89.90) was copied onto seven items.
A cell is emitted once, at its first column and first row; the positions it covers are empty.
"""
from __future__ import annotations

import argparse
import html
import json
import os
import re
import sys
import unicodedata
import zipfile
from datetime import date
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
LEGAL = ROOT / "db/seed/data/legal"
sys.path.insert(0, str(ROOT / "research/legal-loader"))
sys.path.insert(0, str(ROOT / "apps/ingest"))
sys.path.insert(0, str(Path(__file__).parent))
import parse_provisions as pp  # noqa: E402
from build_chunks import build_chunks  # noqa: E402

TODAY = date.today().isoformat()
DOC_TYPE = [("NĐ-CP", "nghi_dinh", "Nghị định"), ("QĐ-", "quyet_dinh", "Quyết định"),
            ("TT-", "thong_tu", "Thông tư"), ("NQ-", "nghi_quyet", "Nghị quyết")]

# (number, congbao id, work sub-directory, domain) — the domain is ours, not Công báo's.
SOURCES = [
    ("292/2026/NĐ-CP", 470149, "292-2026-nd-cp", "Quản lý ngoại thương"),
    ("336/2026/NĐ-CP", 470341, "336-2026-nd-cp", "Thủ tục hải quan"),
    ("37/2026/NĐ-CP", 468865, "37-2026-nd-cp", "Chất lượng sản phẩm, hàng hóa"),
    ("36/2026/TT-BKHCN", 469968, "36-2026-tt-bkhcn", "Danh mục hàng hóa rủi ro"),
    ("72/2022/NĐ-CP", 38012, "72-2022-nd-cp", "In ấn — thiết bị in"),
    ("11/2024/TT-BTTTT", 42827, "11-2024-tt-btttt", "In ấn — thiết bị in"),
    ("52/2018/TT-BCT", 28155, "52-2018-tt-bct", "Năng lượng"),
    ("18/2019/QĐ-TTg", 28889, "18-2019-qd-ttg", "Máy móc, thiết bị đã qua sử dụng"),
    ("33/2026/TT-BCT", 469965, "33-2026-tt-bct", "Danh mục hàng hóa rủi ro"),
    ("41/2026/TT-BXD", 469969, "41-2026-tt-bxd", "Danh mục hàng hóa rủi ro"),
    ("49/2026/TT-BXD", 469975, "49-2026-tt-bxd", "Danh mục hàng hóa rủi ro"),
    ("27/2026/TT-BNNMT", 469989, "27-2026-tt-bnnmt", "Danh mục hàng hóa rủi ro"),
    ("27/2026/TT-BYT", 469998, "27-2026-tt-byt", "Danh mục hàng hóa rủi ro"),
    ("125/2026/TT-BCA", 470259, "125-2026-tt-bca", "Danh mục hàng hóa rủi ro"),
    ("169/2026/NĐ-CP", 469561, "169-2026-nd-cp", "Xử phạt hải quan"),
    ("85/2026/TT-BTC", 469982, "85-2026-tt-btc", "Thủ tục hải quan"),
]


def repack_docx(src: Path) -> Path:
    dst = src.with_name(src.stem + ".fixed.docx")
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename.replace("\\", "/")
            if not name.startswith("[trash]"):
                zout.writestr(name, zin.read(item.filename))
    return dst


from docs_safe import strip_word_field_codes as clean  # noqa: E402  (one definition), decode_symbol_font

DEBRIS = re.compile(r"HYPERLINK|MERGEFORMAT|PAGEREF|_Toc\d")


ANCHOR = re.compile(r"^(Điều\s+\d+\s*\.|PHỤ LỤC|Phụ lục\s+[IVXLC0-9]+\b)")
# A bare "Phụ lục" heading (TT 33/2026/TT-BCT, TT 125/2026/TT-BCA) is read by read_docx only, as
# pp.APPENDIX_HEADING: in the textutil branch it made NĐ 72/2022 flag a table the txt cannot place.
#: Inside an annex, the tier sub-heading that alone tells two tables of one bare "Phụ lục" apart
#: (TT 125/2026/TT-BCA: "A. Danh mục … rủi ro cao", "B. Danh mục … rủi ro trung bình").
SUBHEADING = re.compile(r"^[A-ZĐ]\.\s+Danh mục")
#: An annex headed by nothing but "DANH MỤC" (TT 15/2024/TT-BYT, a legacy .doc: read_doc). Upper case,
#: the whole line, and only after the last Điều — "Danh mục" runs through every article body.
BARE_DANH_MUC = re.compile(r"^DANH MỤC$")
#: Tables that are not annex content: the letterhead (before the first Điều) and the signature
#: block (after the last Điều, before any annex).
LETTERHEAD = re.compile(r"CỘNG H(?:ÒA|OÀ) XÃ HỘI CHỦ NGHĨA VIỆT NAM")
SIGNATURE = re.compile(r"^(TM\.|KT\.|Q\.|TL\.|Nơi nhận|THỦ TƯỚNG\b|BỘ TRƯỞNG\b|CHỦ TỊCH\b)")
#: The caption Công báo gives each gazette issue a document is in: "Công báo số 406 ngày 2026-07-17".
#: Not the _signed.pdf path — its date is the upload day (NĐ 292/2026: /2026/7/31/, gazette 2026-08-03).
GAZETTE = re.compile(r"Công báo số\s+([\d\s+]+?)\s+ngày\s+(\d{4}-\d{2}-\d{2})")


#: Glyphs a Word run draws from a PICTURE font, stored as private-use code points. Google Docs
#: deletes private-use characters, so they must become the Unicode they draw — and the
#: meaning depends on the FONT, not the code point: U+F02A is `∗` in Symbol but the empty
#: tick box of a form in Wingdings 2. Wingdings 2 entries were read off the rendered Công báo
#: PDF of 37/2026/NĐ-CP (form in the GS1 annex: "Điền ✓ vào ô trống" beside □ boxes), not
#: from memory. A code point not listed stays as it is, and verify_drive.py flags it.
_PICTURE_FONTS = {"Wingdings 2": {0xF020: " ", 0xF02A: "□", 0xF050: "✓"}}


def decode_glyph_runs(body, qn) -> None:
    """Rewrite, in place, every run whose font is Symbol or a listed picture font."""
    for run in body.iter(qn("w:r")):
        fonts = run.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        if fonts is None:
            continue
        name = fonts.get(qn("w:hAnsi")) or fonts.get(qn("w:ascii")) or ""
        table = _PICTURE_FONTS.get(name)
        for t in run.iter(qn("w:t")):
            if t.text and any(0xF000 <= ord(c) <= 0xF0FF for c in t.text):
                t.text = t.text.translate(table) if table else (
                    decode_symbol_font(t.text) if "Symbol" in name else t.text)


def restore_auto_numbers(body, numbering, qn) -> None:
    """Write Word's decimal auto-number into each numbered paragraph, in place, as its first run.

    A khoản typed as a Word numbered list (w:numPr) keeps its "1." in numbering.xml, not in the
    text: python-docx reads " Phạt tiền từ …" and the parser sees no khoản 1 (169/2026/NĐ-CP Điều 9,
    85/2026/TT-BTC Điều 1). The number is counted as Word does — per w:num, per level, from the
    level's w:start, a deeper level restarting when a shallower one advances — and printed through
    its w:lvlText ("%1."). Only numFmt decimal is written: bullets and other formats stay as read.
    ponytail: direct w:numPr only (no style-inherited numbering, no w:lvlOverride) — none of the
    Công báo files read so far use either; resolve them here when one does."""
    from docx.oxml import OxmlElement
    if numbering is None:
        return
    val = lambda el, tag: None if el is None or el.find(qn(tag)) is None else el.find(qn(tag)).get(qn("w:val"))
    abstract = {a.get(qn("w:abstractNumId")): a for a in numbering.iter(qn("w:abstractNum"))}
    levels = {}  # (numId, ilvl) → (numFmt, lvlText, start)
    for num in numbering.iter(qn("w:num")):
        for lvl in abstract.get(val(num, "w:abstractNumId"), []):
            if lvl.tag == qn("w:lvl"):
                levels[(num.get(qn("w:numId")), int(lvl.get(qn("w:ilvl"))))] = (
                    val(lvl, "w:numFmt"), val(lvl, "w:lvlText"), int(val(lvl, "w:start") or 1))
    counters = {}  # numId → {ilvl: current number}
    for p in body.iter(qn("w:p")):
        numpr = p.find(qn("w:pPr") + "/" + qn("w:numPr"))
        num_id, ilvl = val(numpr, "w:numId"), int(val(numpr, "w:ilvl") or 0)
        if (num_id, ilvl) not in levels:  # no numPr, or numId 0 = numbering switched off
            continue
        fmt, text, start = levels[(num_id, ilvl)]
        count = counters.setdefault(num_id, {})
        count[ilvl] = count.get(ilvl, start - 1) + 1
        for deeper in [k for k in count if k > ilvl]:
            del count[deeper]
        if fmt != "decimal" or not text:
            continue
        label = re.sub(r"%(\d)", lambda m: str(count.get(int(m.group(1)) - 1, 1)), text)
        typed = "".join(t.text or "" for t in p.iter(qn("w:t")))
        run, t = OxmlElement("w:r"), OxmlElement("w:t")
        t.text = label if typed[:1].isspace() else label + " "  # " Phạt" → "1. Phạt", not "1.  Phạt"
        run.append(t)
        p.find(qn("w:pPr")).addnext(run)


def grid_cells(tr, table) -> list[str]:
    """One text per grid column of a `w:tr`: a merged cell at its first column (gridSpan) and
    first row (vMerge restart) only, so a row keeps its width and nothing is read twice."""
    from docx.table import _Cell
    out = []
    for tc in tr.tc_lst:
        out.append("" if tc.vMerge == "continue" else _Cell(tc, table).text.strip())
        out += [""] * (tc.grid_span - 1)
    return out


def read_docx(path: Path) -> tuple[list[str], list[dict | None]]:
    """Prose lines, and every table under the last heading before it.

    Once the annexes begin, the heading is the annex's: a "Điều" of a form template (NĐ 37/2026
    Phụ lục VII) no longer moves it, a tier sub-heading is appended to it ("Phụ lục — A. Danh mục
    …"), and an annex title merged across row 1 of its own table (TT 49/2026/TT-BXD Phụ lục I)
    becomes it and leaves the rows. Prose after an annex table ("Ghi chú", footnotes) is kept with
    that table as `notes`. The letterhead and the signature block come back as None, so every other
    table keeps its place in the document as its index."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(repack_docx(path)))
    decode_glyph_runs(d.element.body, qn)
    numbering = d.part.part_related_by(RT.NUMBERING).element if any(
        r.reltype == RT.NUMBERING for r in d.part.rels.values()) else None
    restore_auto_numbers(d.element.body, numbering, qn)
    # Matched in NFC, kept as typed: TT 36/2026's letterhead spells "CỘNG" with a combining dot.
    nfc = lambda s: unicodedata.normalize("NFC", s)
    heading = lambda s: ANCHOR.match(s) or pp.APPENDIX_HEADING.match(s)
    lines, tables, anchor, annex, notes = [], [], None, None, None
    for el in d.element.body.iterchildren():
        if el.tag == qn("w:p"):
            for line in Paragraph(el, d._body).text.split("\n"):
                line, key = line.strip(), nfc(line.strip())
                if not line:
                    continue
                lines.append(line)
                if heading(key) and not (annex and key.startswith("Điều")):
                    anchor, notes = line[:120], None
                    if not key.startswith("Điều"):
                        annex = anchor
                elif annex and SUBHEADING.match(key):
                    anchor, notes = f"{annex} — {line[:120]}", None
                elif notes is not None:
                    notes.append(line)
        elif el.tag == qn("w:tbl"):
            table = Table(el, d._body)
            rows = [grid_cells(tr, table) for tr in el.tr_lst]
            first = el.tr_lst[0].tc_lst if rows else []
            title = rows[0][0].split("\n")[0].strip() if len(first) == 1 and first[0].grid_span > 1 else ""
            if heading(nfc(title)) and not nfc(title).startswith("Điều"):
                annex = anchor = title[:120]
                rows, notes = rows[1:], None
            cells = [nfc(c) for r in rows for c in r]
            if annex is None and (any(LETTERHEAD.search(c) for c in cells) if anchor is None
                                  else any(SIGNATURE.match(c) for c in cells)):
                tables.append(None)
                continue
            tables.append({"anchor": anchor, "rows": rows})
            if annex:
                notes = tables[-1]["notes"] = []
    return lines, tables


class HtmlTables(HTMLParser):
    """Top-level tables of textutil's HTML export as rows → cells → paragraph texts, each
    flagged with whether an annex heading ("Phụ lục I", "PHỤ LỤC") came before it, and with
    how many "Điều" had passed when the last bare "DANH MỤC" before it was seen (`bare`):
    it is an annex heading only if that equals `dieu`, the document's total."""

    def __init__(self):
        super().__init__()
        self.tables, self.depth, self.para, self.annex = [], 0, None, False
        self.dieu, self.bare = 0, None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self.depth += 1
            if self.depth == 1:
                self.tables.append({"annex": self.annex, "bare": self.bare, "rows": []})
        elif tag == "tr" and self.depth == 1:
            self.tables[-1]["rows"].append([])
        elif tag == "td" and self.depth == 1:
            self.tables[-1]["rows"][-1].append([])
        elif tag == "p":
            self.para = []

    def handle_endtag(self, tag):
        if tag == "table":
            self.depth -= 1
        elif tag == "p" and self.para is not None:
            text = "".join(self.para).strip()
            if self.depth:
                self.tables[-1]["rows"][-1][-1].append(text)
            elif ANCHOR.match(text) and not text.startswith("Điều"):
                self.annex = True
            else:
                key = unicodedata.normalize("NFC", text)
                if key.startswith("Điều") and ANCHOR.match(key):
                    self.dieu += 1
                elif BARE_DANH_MUC.match(key):
                    self.bare = self.dieu
            self.para = None

    def handle_data(self, data):
        if self.para is not None:
            self.para.append(data)


def read_doc(path: Path) -> tuple[list[str], list[dict]]:
    """Legacy .doc via textutil: prose lines through the project reader, table rows
    rebuilt from the Word cell mark (\\x07), exactly as the tariff loaders do.

    A table textutil RECOGNISES loses its cell marks in the txt export: each cell paragraph
    comes out as a bare line. The cell-mark branch never saw those, and the prose parser stops
    at the annex heading (R9), so QĐ 18/2019 lost its whole Phụ lục I (HS heading → maximum
    equipment age) between the two. The HTML export keeps such a table's shape; the table is
    placed where its paragraphs run in the txt, and its cells take the txt lines verbatim.
    Only tables after an annex heading are taken — before it they are the letterhead and the
    signature block. A bare "DANH MỤC" counts as one only after the last Điều (TT 15/2024/TT-BYT).
    A located table must match line for line, or the document is refused."""
    import subprocess
    run = lambda fmt: subprocess.run(["textutil", "-convert", fmt, "-stdout", str(path)],
                                     capture_output=True, text=True).stdout
    html_tables = HtmlTables()
    html_tables.feed(run("html"))
    pending = [t["rows"] for t in html_tables.tables if t["annex"] or t["bare"] == html_tables.dieu]
    lines = [clean(l) for l in run("txt").split("\n")]
    key = lambda l: unicodedata.normalize("NFC", l.strip())
    last_dieu = max((n for n, l in enumerate(lines)
                     if "\x07" not in l and key(l).startswith("Điều") and ANCHOR.match(key(l))), default=-1)
    tables, current, anchor, i = [], None, None, 0
    while i < len(lines):
        line = lines[i]
        if "\x07" in line:
            cells = [c.strip() for c in line.split("\x07")]
            while cells and not cells[-1]:
                cells.pop()
            if current is None:
                current = {"anchor": anchor, "rows": []}
                tables.append(current)
            if cells:
                current["rows"].append(cells)
        else:
            current = None
            paras = [p for row in pending[0] for cell in row for p in cell] if pending else []
            if paras and anchor and not anchor.startswith("Điều") and \
                    [l.split() for l in lines[i:i + len(paras)]] == [p.split() for p in paras]:
                rows = []
                for row in pending.pop(0):
                    cells = []
                    for cell in row:
                        cells.append("\n".join(l.strip() for l in lines[i:i + len(cell)]).strip())
                        i += len(cell)
                    while cells and not cells[-1]:
                        cells.pop()
                    if cells:
                        rows.append(cells)
                tables.append({"anchor": anchor, "rows": rows})
                continue
            if ANCHOR.match(line.strip()) or (i > last_dieu and BARE_DANH_MUC.match(key(line))):
                anchor = line.strip()[:120]
        i += 1
    if pending:
        raise SystemExit(f"{path}: {len(pending)} bảng phụ lục (HTML) không khớp dòng nào của bản txt — không ghi")
    return [clean(l) for l in pp.doc_lines(str(path))], tables


def content_page(cid: int) -> str:
    import ingest_document as ig
    return ig.get(f"https://congbao.chinhphu.vn/van-ban/x-{cid}.htm").decode("utf-8", "ignore")


def page_title(page: str) -> str:
    m = re.search(r"<title>(.*?)</title>", page, re.S)
    return re.sub(r"\s+", " ", html.unescape(m.group(1))).strip() if m else ""


def gazette(page: str) -> tuple[str | None, str | None]:
    """(gazette_issue, gazette_date) from the page's issue captions: "99 &#x2B; 100" → "99+100";
    a document spread over several issues gives "first…last" and the first issue's date."""
    found = sorted({(d, re.sub(r"\s+", "", n)) for n, d in GAZETTE.findall(html.unescape(page))},
                   key=lambda x: (x[0], int(re.match(r"\d+", x[1]).group())))
    if not found:
        return None, None
    (date_, first), last = found[0], found[-1][1]
    return (first if first == last else f"{first}…{last}"), date_


def doc_row(number: str, cid: int, domain: str) -> dict:
    import ingest_document as ig
    meta = ig.fetch_metadata(cid)
    code, label = next((c, l) for m, c, l in DOC_TYPE if m in number)
    page = content_page(cid)
    issue, gazette_date = gazette(page)
    title = page_title(page)
    title = re.sub(rf"^{label}\s+số\s+{re.escape(number)}\s*", "", title, flags=re.I).strip(" .-–|")
    title = re.split(r"\s+[-|]\s+Công báo", title)[0]
    title = title[:1].upper() + title[1:]   # hand-built rows are capitalised; the summary
    # reads as a sentence ("Nghị định 31/2018/NĐ-CP quy định…") and keeps it lower — as they do
    eff = meta.get("effective_from")
    if not eff:
        raise SystemExit(f"{number}: Công báo không nêu ngày hiệu lực — từ chối, không đoán")
    return {
        "number": number, "doc_type": code, "title": title[:500],
        "issuing_body": meta.get("issuing_body"), "signed_date": meta.get("signed_date"),
        "gazette_date": meta.get("gazette_date") or gazette_date, "effective_from": eff, "effective_to": None,
        # Công báo publishes no machine-readable in-force status (it reports NĐ 134/2016
        # as never amended). The only thing we can state is whether the effective date has
        # arrived; everything else is carried by `auto_unverified` and its warning.
        "effectiveness": "chua_co_hieu_luc" if eff > TODAY else "con_hieu_luc",
        "is_consolidated": False, "consolidates": None, "gazette_issue": issue,
        "source_url": f"https://congbao.chinhphu.vn/van-ban/x-{cid}.htm",
        "summary": f"{label} {number} {title[:1].lower() + title[1:]}"[:400], "domain": domain,
        "short": f"{label} {number}", "verification": "auto_unverified", "verified_by": None,
    }


def download(number: str, cid: int, target: Path) -> None:
    """Fetch every part (.doc/.docx/.pdf) of one document from Công báo into `target`.

    curl, not urllib: the download CDN sends only its leaf certificate and omits the
    GlobalSign intermediate. macOS curl verifies through the system trust store; Python's
    OpenSSL does not, and fails with CERTIFICATE_VERIFY_FAILED. (The Docker image solves
    the same thing by pinning the intermediate — apps/ingest/Dockerfile.) The g7 links are
    tokenised and expire fast, so they are read from the page and used at once.
    """
    import subprocess
    import ingest_document as ig
    if target.exists() and any(target.iterdir()):
        return
    target.mkdir(parents=True, exist_ok=True)
    parts = ig.fetch_part_urls(f"https://congbao.chinhphu.vn/van-ban/x-{cid}.htm")
    if not parts:
        raise SystemExit(f"{number}: trang Công báo không có file tải về")
    for i, (name, url) in enumerate(parts, 1):
        out = target / f"{i:02d}{Path(name).suffix.lower() or '.bin'}"
        r = subprocess.run(["curl", "-sS", "-L", "--max-time", "240", "-o", str(out), "-w", "%{http_code}", url],
                           capture_output=True, text=True)
        if r.stdout != "200":
            raise SystemExit(f"{number}: tải phần {i} lỗi HTTP {r.stdout} {r.stderr[:120]}")


def gate(number: str, rows: list[dict]) -> None:
    nums = [int(re.match(r"\d+", str(r["number"])).group()) for r in rows if r["ptype"] == "dieu"]
    if not nums or nums != list(range(1, len(nums) + 1)):
        missing = sorted(set(range(1, max(nums or [0]) + 1)) - set(nums))
        raise SystemExit(f"{number}: CỔNG CẤU TRÚC — Điều không liên tục (thiếu {missing[:10]}). Không ghi.")


def gate_debris(number: str, rows: list[dict]) -> None:
    dirty = [r["key"] for r in rows if DEBRIS.search((r.get("body") or "") + (r.get("heading") or ""))]
    if dirty:
        raise SystemExit(f"{number}: CỔNG RÁC MÃ TRƯỜNG — {len(dirty)} điều khoản còn lệnh trường Word ({dirty[:3]}). Không ghi.")


def load(path: Path) -> list[dict]:
    return [json.loads(l) for l in path.open(encoding="utf-8")] if path.exists() else []


def save(path: Path, rows: list[dict]) -> None:
    with path.open("w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(json.dumps(row, ensure_ascii=False) + "\n")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--work", required=True)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--download", action="store_true", help="tải từ Công báo văn bản nào chưa có trong --work")
    ap.add_argument("--only", action="append", metavar="SỐ_HIỆU",
                    help="chỉ nạp lại văn bản này (lặp lại được); dòng của văn bản khác giữ nguyên")
    ap.add_argument("--out", metavar="THƯ_MỤC",
                    help="ghi bốn file vào thư mục này thay cho db/seed/data/legal/ (xem trước trong thư mục tạm)")
    args = ap.parse_args()
    work, out = Path(args.work), Path(args.out or LEGAL)
    unknown = set(args.only or []) - {s[0] for s in SOURCES}
    if unknown:
        raise SystemExit(f"--only: không có trong SOURCES: {sorted(unknown)}")

    new_docs, new_provs, new_tables = {}, [], []
    for number, cid, stem, domain in SOURCES:
        if args.only and number not in args.only:
            continue
        d = work / stem
        if args.download:
            download(number, cid, d)
        src = next((p for p in sorted(d.glob("*.docx")) if ".fixed." not in p.name), None) \
            or next(iter(sorted(d.glob("*.doc"))), None)
        if src is None:
            raise SystemExit(f"{number}: không có file Word trong {d}")
        lines, tables = read_docx(src) if src.suffix == ".docx" else read_doc(src)
        doc = doc_row(number, cid, domain)
        pp.read_doc = lambda _doc, _l=lines: _l  # reuse the project's article logic verbatim
        rows = pp.build(doc)
        gate(number, rows)
        gate_debris(number, rows)
        new_docs[number] = doc
        new_provs += rows
        tables = [(i, t) for i, t in enumerate(tables, 1) if t]  # index = place in the document
        for i, t in tables:
            new_tables.append({"document_number": number, "index": i, "anchor": t["anchor"],
                               "header": t["rows"][0] if t["rows"] else [], "rows": t["rows"][1:],
                               **({"notes": t["notes"]} if t.get("notes") else {})})
        kinds = {k: sum(1 for r in rows if r["ptype"] == k) for k in ("chuong", "dieu", "khoan", "diem")}
        print(f"✅ {number:18} {src.suffix[1:]:4} {kinds} · {len(tables)} bảng/{sum(len(t['rows']) for _, t in tables)} dòng"
              f" · hiệu lực {doc['effective_from']} ({doc['effectiveness']})")

    chunks = build_chunks(new_docs, new_provs)
    print(f"   {len(chunks)} chunk (một chunk mỗi Điều)")
    if args.dry_run:
        return 0
    out.mkdir(parents=True, exist_ok=True)
    keep = lambda name, key: [r for r in load(out / name) if r[key] not in new_docs]
    save(out / "documents.ndjson", keep("documents.ndjson", "number") + list(new_docs.values()))
    save(out / "provisions.ndjson", keep("provisions.ndjson", "document_number") + new_provs)
    save(out / "chunks.ndjson", keep("chunks.ndjson", "document_number") + chunks)
    save(out / "annex-tables.ndjson", keep("annex-tables.ndjson", "document_number") + new_tables)
    print(f"Đã ghi {out}: {len(load(out / 'documents.ndjson'))} văn bản · {len(load(out / 'provisions.ndjson')):,} điều khoản · "
          f"{len(load(out / 'chunks.ndjson')):,} chunk · {len(load(out / 'annex-tables.ndjson'))} bảng phụ lục")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
