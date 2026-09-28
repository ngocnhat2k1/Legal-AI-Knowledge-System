"""Owner's digest sheet of classification rulings (PTPL) -> db/seed/data/legal/classification-digest.ndjson.

The sheet is a third party's summary: code, trade name, short description, deciding feature, number, link. It is
not the rulings' text, so seed-evidence files each row as a `reference` ruling labelled as a summary.

    .venv/bin/python research/inbox-loader/ptpl_digest.py <sheet URL | sheet.xlsx>

The sheet URL is kept out of this public repo, in .agent/local/ptpl-digest.md.

Rerun when the owner adds rows; seed-evidence embeds only the new ones.
"""
import io
import json
import re
import sys
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'db/seed/data/legal/classification-digest.ndjson'
COLUMNS = ['stt', 'nam', 'chuong', 'nhom_hang', 'ma_hs', None, 'ten_thuong_mai', 'mo_ta_ma_hs', 'mo_ta_hang',
           'dac_tinh', 'so_hieu', 'co_quan', 'lien_ket']


def link(cell):
    """=HYPERLINK("url","label") -> url; a plain value passes through."""
    m = re.match(r'=HYPERLINK\("([^"]+)"', str(cell or ''))
    return m.group(1) if m else cell


def rows(book):
    ws = book['PTPL ma HS']
    header = [c.value for c in ws[4]][1:]
    assert header[0] == 'STT' and header[4] == 'Mã HS' and header[10] == 'Số hiệu văn bản', header
    ahtn = {json.loads(l)['hs'] for l in (ROOT / 'db/seed/data/hs-descriptions.ndjson').open()}
    for values in ws.iter_rows(min_row=5, values_only=True):
        if values[1] is None:
            continue
        r = {k: (v.strip() if isinstance(v, str) else v) for k, v in zip(COLUMNS, values[1:]) if k}
        r['lien_ket'] = link(r['lien_ket'])
        r['ma_hs'] = str(r['ma_hs'])
        assert re.fullmatch(r'\d{8}', r['ma_hs']), r
        r['ma_trong_ahtn_2022'] = r['ma_hs'] in ahtn
        yield r


def main():
    src = sys.argv[1]
    if src.startswith('https://'):  # a Google Sheets link; the sheet must be shared by link
        src = io.BytesIO(urllib.request.urlopen(re.sub(r'/edit.*$', '', src) + '/export?format=xlsx').read())
    book = openpyxl.load_workbook(src)
    out = list(rows(book))
    OUT.write_text(''.join(json.dumps(r, ensure_ascii=False) + '\n' for r in out))
    stale = [r['ma_hs'] for r in out if not r['ma_trong_ahtn_2022']]
    print(f'{len(out)} rows -> {OUT.relative_to(ROOT)}; not in AHTN 2022: {stale or "none"}')


if __name__ == '__main__':
    main()
