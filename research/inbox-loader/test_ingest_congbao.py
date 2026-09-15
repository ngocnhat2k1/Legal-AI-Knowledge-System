"""Tests for the annex-table reader of ingest_congbao.py and the gazette caption.

Each case is a shape found in a Công báo risk-list circular (2026-09-14): a group name spread
over the HS column (TT 125/2026/TT-BCA), one merged cell of 13 codes over seven items (same), a
bare "Phụ lục" split into tier tables A/B (same), an annex title inside row 1 (TT 49/2026/TT-BXD);
a khoản 1 typed as a Word numbered list (169/2026/NĐ-CP Điều 9), and an annex headed by a bare
"DANH MỤC" in a legacy .doc (TT 15/2024/TT-BYT).

    .venv/bin/python -m unittest test_ingest_congbao     (needs python-docx; the .doc case needs macOS textutil)
"""
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from ingest_congbao import gazette, read_doc, read_docx

CODES = "8525.81.10\n8525.81.20; 8525.89.90"


def read(fill, part=1):
    with tempfile.TemporaryDirectory() as tmp:
        d = docx.Document()
        fill(d)
        path = Path(tmp) / "t.docx"
        d.save(path)
        return read_docx(path)[part]


def header(table, names):
    for i, name in enumerate(names):
        table.cell(0, i).text = name


class TestCells(unittest.TestCase):
    def test_grid_span_does_not_shift_or_repeat(self):
        def fill(d):
            d.add_paragraph("Phụ lục")
            t = d.add_table(rows=2, cols=4)
            header(t, ["TT", "Tên", "Mã HS", "QCVN"])
            t.cell(1, 0).text = "I"
            t.cell(1, 1).merge(t.cell(1, 3)).text = "Pháo hoa"
        [table] = read(fill)
        self.assertEqual(table["rows"][1], ["I", "Pháo hoa", "", ""])

    def test_vertical_merge_is_read_once(self):
        def fill(d):
            d.add_paragraph("Phụ lục")
            t = d.add_table(rows=4, cols=3)
            header(t, ["TT", "Tên", "Mã HS"])
            for r in (1, 2, 3):
                t.cell(r, 0).text, t.cell(r, 1).text = str(r), f"Camera {r}"
            t.cell(1, 2).merge(t.cell(3, 2)).text = CODES
        [table] = read(fill)
        self.assertEqual([r[2] for r in table["rows"][1:]], [CODES, "", ""])
        self.assertEqual([r[1] for r in table["rows"][1:]], ["Camera 1", "Camera 2", "Camera 3"])

    def test_multi_code_cell_keeps_raw_text(self):
        def fill(d):
            d.add_paragraph("Phụ lục I")
            t = d.add_table(rows=2, cols=2)
            header(t, ["Tên", "Mã HS"])
            t.cell(1, 0).text, t.cell(1, 1).text = "Camera", CODES
        [table] = read(fill)
        self.assertEqual(table["rows"][1], ["Camera", CODES])


class TestAnchors(unittest.TestCase):
    def test_bare_phu_luc_and_tier_subheadings(self):
        def fill(d):
            d.add_paragraph("Điều 7. Trách nhiệm thi hành")
            d.add_paragraph("Phụ lục")
            d.add_paragraph("DANH MỤC SẢN PHẨM, HÀNG HOÁ CÓ MỨC ĐỘ RỦI RO TRUNG BÌNH, MỨC ĐỘ RỦI RO CAO")
            d.add_paragraph("A. Danh mục sản phẩm, hàng hoá có mức độ rủi ro cao")
            d.add_table(rows=1, cols=2).cell(0, 0).text = "TT"
            d.add_paragraph("B. Danh mục sản phẩm, hàng hoá có mức độ rủi ro trung bình")
            d.add_table(rows=1, cols=2).cell(0, 0).text = "TT"
        self.assertEqual([t["anchor"] for t in read(fill)], [
            "Phụ lục — A. Danh mục sản phẩm, hàng hoá có mức độ rủi ro cao",
            "Phụ lục — B. Danh mục sản phẩm, hàng hoá có mức độ rủi ro trung bình",
        ])

    def test_annex_title_in_row_one(self):
        def fill(d):
            d.add_paragraph("Điều 7. Điều khoản chuyển tiếp")
            t = d.add_table(rows=2, cols=3)
            t.cell(0, 0).merge(t.cell(0, 2)).text = "Phụ lục I\nDANH MỤC SẢN PHẨM, HÀNG HÓA CÓ MỨC ĐỘ RỦI RO CAO"
            for i, name in enumerate(["STT", "Tên", "Mã số HS"]):
                t.cell(1, i).text = name
            d.add_table(rows=1, cols=1).cell(0, 0).text = "tiếp"
        tables = read(fill)
        self.assertEqual([t["anchor"] for t in tables], ["Phụ lục I", "Phụ lục I"])
        self.assertEqual(tables[0]["rows"], [["STT", "Tên", "Mã số HS"]])

    def test_letterhead_signature_dropped_annex_prose_kept(self):
        def fill(d):
            d.add_table(rows=1, cols=2).cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc"
            d.add_paragraph("Điều 1. Hiệu lực")
            d.add_table(rows=1, cols=2).cell(0, 1).text = "KT. BỘ TRƯỞNG\nTHỨ TRƯỞNG"
            d.add_paragraph("Phụ lục I")
            d.add_table(rows=1, cols=2).cell(0, 0).text = "TT"
            d.add_paragraph("Ghi chú:")
            d.add_paragraph("- Nguyên liệu dùng chung cho thức ăn chăn nuôi và thủy sản")
            d.add_paragraph("Phụ lục II")
            d.add_paragraph("Điều 3. ..... (điều của một mẫu quyết định)")
            d.add_table(rows=1, cols=2).cell(0, 0).text = "Nơi nhận:"
        tables = read(fill)
        self.assertEqual(tables[:2], [None, None])
        self.assertEqual(tables[2]["notes"], ["Ghi chú:", "- Nguyên liệu dùng chung cho thức ăn chăn nuôi và thủy sản"])
        self.assertEqual(tables[3]["anchor"], "Phụ lục II")


class TestAutoNumber(unittest.TestCase):
    def test_word_numbered_khoan_gets_its_number(self):
        def fill(d):
            numbering = d.part.numbering_part.element
            for num_id, fmt, text in ((90, "decimal", "%1."), (91, "bullet", "-")):
                numbering.append(parse_xml(
                    f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{num_id}"><w:lvl w:ilvl="0"><w:start w:val="1"/>'
                    f'<w:numFmt w:val="{fmt}"/><w:lvlText w:val="{text}"/></w:lvl></w:abstractNum>'))
                numbering.append(parse_xml(
                    f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{num_id}"/></w:num>'))
            d.add_paragraph("Điều 9. Vi phạm quy định về khai hải quan")
            for text, num_id in ((" Phạt tiền đối với một trong các hành vi vi phạm sau:", 90), ("a) Khai sai", None),
                                 ("Phạt tiền đối với hành vi khác", 90), ("…………", 91), ("Không đánh số", 0)):
                p = d.add_paragraph(text)
                if num_id is not None:
                    p._p.get_or_add_pPr().insert(0, parse_xml(
                        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="0"/><w:numId w:val="{num_id}"/></w:numPr>'))
        self.assertEqual(read(fill, part=0)[1:], [
            "1. Phạt tiền đối với một trong các hành vi vi phạm sau:", "a) Khai sai",
            "2. Phạt tiền đối với hành vi khác", "…………", "Không đánh số"])


@unittest.skipUnless(shutil.which("textutil"), "read_doc reads a .doc through macOS textutil")
class TestDocAnnex(unittest.TestCase):
    def test_bare_danh_muc_anchors_only_after_the_last_dieu(self):
        page = ('<html><head><meta charset="utf-8"></head><body>'
                "<p>Điều 1. Ban hành Danh mục</p><p>DANH MỤC</p>"
                '<table border="1"><tr><td><p>A</p></td><td><p>B</p></td></tr></table>'
                "<p>Điều 2. Hiệu lực thi hành</p><p>KT. BỘ TRƯỞNG</p><p>DANH MỤC</p><p>Thực phẩm</p>"
                '<table border="1"><tr><td><p>STT</p></td><td><p>Mã HS</p></td></tr>'
                "<tr><td><p>1</p></td><td><p>1518.00.14</p></td></tr></table></body></html>")
        with tempfile.TemporaryDirectory() as tmp:
            src, doc = Path(tmp) / "t.html", Path(tmp) / "t.doc"
            src.write_text(page, encoding="utf-8")
            subprocess.run(["textutil", "-convert", "doc", str(src), "-output", str(doc)], check=True)
            tables = read_doc(doc)[1]
        self.assertEqual(tables, [{"anchor": "DANH MỤC", "rows": [["STT", "Mã HS"], ["1", "1518.00.14"]]}])


class TestGazette(unittest.TestCase):
    def test_issue_and_date_from_caption(self):
        self.assertEqual(gazette('<a title="Công báo số 405 ngày 2026-07-17">405</a>'), ("405", "2026-07-17"))
        self.assertEqual(gazette('title="Công báo số 99 &#x2B; 100 ngày 2019-01-28"'), ("99+100", "2019-01-28"))
        self.assertEqual(gazette("<html></html>"), (None, None))


if __name__ == "__main__":
    unittest.main()
